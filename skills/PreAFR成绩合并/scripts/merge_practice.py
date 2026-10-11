# -*- coding: utf-8 -*-
"""
学测成绩文档四合一 —— 合并脚本(可复用、版本无关)
================================================
用法:
    python merge_practice.py <原始文件夹> <输出文件.xlsx> [--backup 现有文件备份路径] [--score-sheet-name 名称]

原始文件夹中应包含每次练习导出的文件(班级数不限,按班名排序先后拼接):
    * <班名>-英语-答题明细*.xlsx      每班一份(含 学生明细*/客观题填涂/学生错题汇总 等sheet)
    * 其余xlsx(班级小题明细/学生得分明细等) 每班一份,单sheet,含"序号"表头行+纯数字题号列

兼容不同导出版本:
    * 答题明细的sheet名与列结构按表头名动态识别(题号/题型列位置不限)
    * 得分明细的题号列 = "序号"表头行中所有纯数字单元格(位置不限)

输出工作簿:
    1. [得分汇总sheet](默认名"班级小题得分",可用--score-sheet-name改) ——
       各班得分明细纵向拼接,题号表头按答题明细的题型改写(如 阅读理解1/信息匹配16),
       改写后的换行表头行高设为28.8磅
    2~N. 答题明细的各个sheet —— 各班同名sheet纵向拼接(每班保留自己的标题行+表头行)

样式(字体/边框/填充/行高/列宽/合并单元格)全部从原始文件保真复制;内置逐格校验。
"""
import os
import re
import sys
import shutil
import argparse
from copy import copy
from openpyxl import load_workbook

CLASS_RE = re.compile(r"(高三\d+班[^-]*)")
NUM_RE = re.compile(r"^\d+(\.0)?$")


def as_int(v):
    """单元格值→整数(纯数字才行),否则None"""
    if v is None:
        return None
    s = str(v).strip()
    if not NUM_RE.match(s):
        return None
    return int(float(s))


def find_class_files(raw_dir):
    """按班名配对 答题明细 + 得分明细文件,返回按班名排序的列表"""
    pairs = {}
    for fn in os.listdir(raw_dir):
        if not fn.lower().endswith(".xlsx") or fn.startswith("~$"):
            continue
        m = CLASS_RE.search(fn)
        if not m:
            continue
        cls = m.group(1)
        d = pairs.setdefault(cls, {})
        if "答题明细" in fn:
            d["detail"] = os.path.join(raw_dir, fn)
        else:
            d.setdefault("score", os.path.join(raw_dir, fn))
    classes = sorted(pairs.keys())
    missing = [(c, k) for c in classes for k in ("detail", "score") if k not in pairs[c]]
    if missing:
        raise SystemExit(f"错误: 以下班级缺少文件: {missing}")
    if not classes:
        raise SystemExit(f"错误: 在 {raw_dir} 中未找到含班级名(高三N班)的xlsx文件")
    return [(c, pairs[c]["detail"], pairs[c]["score"]) for c in classes]


def detail_sheets(detail_path):
    """答题明细工作簿的sheet名列表(保持原顺序)"""
    wb = load_workbook(detail_path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def extract_type_map(detail_path):
    """从第一个"学生明细"sheet按表头名定位 题号/题型 列,提取 题号→题型 映射"""
    wb = load_workbook(detail_path, data_only=True)
    sname = next((s for s in wb.sheetnames if s.strip().startswith("学生明细")), None)
    if sname is None:
        raise SystemExit(f"错误: {detail_path} 中找不到'学生明细'sheet")
    ws = wb[sname]
    # 表头行 = 含"题号"和"题型"的那一行(从前5行找);数据从表头行的下一行起
    # (新版明细每题占两行[数据行+合并空行],题1在表头后第1行——旧写法固定从第6行起会漏掉前两题)
    rows = list(ws.iter_rows(min_row=1, max_row=5))
    qcol = tcol = None
    data_start = 3
    for i, row in enumerate(rows):
        vals = [str(c.value).strip() if c.value is not None else "" for c in row]
        if "题号" in vals and "题型" in vals:
            qcol = vals.index("题号") + 1
            tcol = vals.index("题型") + 1
            data_start = i + 2
            break
    if qcol is None:
        raise SystemExit(f"错误: {detail_path} 的学生明细中找不到 题号/题型 表头")
    tmap = {}
    for row in ws.iter_rows(min_row=data_start, values_only=True):
        n = as_int(row[qcol - 1] if len(row) >= qcol else None)
        t = row[tcol - 1] if len(row) >= tcol else None
        if n is not None and t:
            tmap.setdefault(n, str(t).strip())
    wb.close()
    if not tmap:
        raise SystemExit(f"错误: {detail_path} 未能提取到 题号→题型 映射")
    # 完整试卷(含66应用文/67读后续写)时,明细对选择题统称"单选题",按标准卷结构细化语义
    if 66 in tmap and 67 in tmap:
        std = {**{n: "听力理解" for n in range(1, 21)},
               **{n: "阅读理解" for n in range(21, 36)},
               **{n: "信息匹配" for n in range(36, 41)},
               **{n: "完形填空" for n in range(41, 56)},
               **{n: "语法填空" for n in range(56, 66)},
               66: "应用文", 67: "读后续写"}
        applied = 0
        for n, t in std.items():
            if n in tmap and tmap[n] != t:
                tmap[n] = t
                applied += 1
        if applied:
            print(f"完整试卷结构: {applied}个题号的题型按标准卷细化(听力/阅读理解/信息匹配/完形填空/语法填空/应用文/读后续写)")
    return tmap


def copy_cell(src_cell, dst_cell):
    """值 + 样式保真复制（_style整体直拷——一次赋值替代6次属性深拷贝，快数倍且保真一致；
    number_format 单独回写以兼容跨工作簿的自定义格式注册表）"""
    dst_cell.value = src_cell.value
    if src_cell.has_style:
        dst_cell._style = copy(src_cell._style)
        dst_cell.number_format = src_cell.number_format


def append_sheet_rows(src_ws, dst_ws, offset, rename_header_map=None, header_row_height=None):
    """把 src_ws 全部行追加到 dst_ws 第 offset+1 行起;合并单元格/行高随偏移平移;
    rename_header_map: {题号int: 新表头文字},作用于"序号"表头行的纯数字单元格(位置不限);
    header_row_height: 改写表头后的行高(换行表头需足够高度)"""
    for row in src_ws.iter_rows():                      # 顺序流式访问，远快于逐格cell()
        for src_cell in row:
            copy_cell(src_cell, dst_ws.cell(row=src_cell.row + offset, column=src_cell.column))
    for r, rd in src_ws.row_dimensions.items():         # 行高按已设维度平移
        if rd.height is not None:
            dst_ws.row_dimensions[r + offset].height = rd.height
    for mr in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(
            start_row=mr.min_row + offset, start_column=mr.min_col,
            end_row=mr.max_row + offset, end_column=mr.max_col,
        )
    for letter, dim in src_ws.column_dimensions.items():
        if dim.width is not None and letter not in dst_ws.column_dimensions:
            dst_ws.column_dimensions[letter].width = dim.width
    # 题号表头改写:"序号"表头行中所有纯数字单元格
    if rename_header_map:
        for r in range(offset + 1, offset + src_ws.max_row + 1):
            if dst_ws.cell(row=r, column=1).value == "序号":
                for c in range(2, dst_ws.max_column + 1):
                    n = as_int(dst_ws.cell(row=r, column=c).value)
                    if n is not None and n in rename_header_map:
                        dst_ws.cell(row=r, column=c).value = rename_header_map[n]
                if header_row_height is not None:
                    dst_ws.row_dimensions[r].height = header_row_height


def build(raw_dir, out_path, score_sheet_name=None):
    classes = find_class_files(raw_dir)
    print("发现班级(按班名排序):")
    for c, d, s in classes:
        print(f"  {c}: {os.path.basename(d)} | {os.path.basename(s)}")

    # 题号→题型映射(各班分别提取并互相校验)
    maps = [extract_type_map(d) for _, d, _ in classes]
    tmap = maps[0]
    for i, m in enumerate(maps[1:], 1):
        if m != tmap:
            diff = {k for k in set(tmap) | set(m) if tmap.get(k) != m.get(k)}
            print(f"警告: 第{i+1}个班题型映射不一致(题号{sorted(diff)}),以第一个班为准")
    print(f"题型映射({len(tmap)}题): "
          f"{min(tmap)}→{tmap[min(tmap)]} ... {max(tmap)}→{tmap[max(tmap)]}")
    rename_map = {n: f"{t}{n}" for n, t in tmap.items()}

    # 以第一个班的答题明细工作簿为基础(保留其原生主题与样式)
    base = load_workbook(classes[0][1])
    base_sheets = base.sheetnames
    for cls, detail, _ in classes[1:]:
        wb = load_workbook(detail)
        if wb.sheetnames != base_sheets:
            print(f"警告: {cls} 答题明细sheet与首班不一致 "
                  f"({wb.sheetnames} vs {base_sheets}),按首班顺序尽力匹配")
        for sname in base_sheets:
            if sname in wb.sheetnames:
                append_sheet_rows(wb[sname], base[sname], offset=base[sname].max_row)
        wb.close()
        print(f"已追加 {cls} 的{len(base_sheets)}个答题明细sheet")

    # 得分汇总sheet放最前:拼接各班得分明细,并改写题号表头
    src0 = load_workbook(classes[0][2])
    src0_name = src0.sheetnames[0]
    if score_sheet_name is None:
        score_sheet_name = "班级小题得分" if "小题明细" in src0_name else src0_name
    ws = base.create_sheet(score_sheet_name, 0)
    offset = 0
    src0.close()
    for cls, _, score in classes:
        swb = load_workbook(score)
        sws = swb.worksheets[0]
        append_sheet_rows(
            sws, ws, offset=offset,
            rename_header_map=rename_map, header_row_height=28.8,
        )
        offset += sws.max_row
        swb.close()
        print(f"已拼接 {cls} 的{src0_name}")
    base.properties.creator = "Z.ai"
    base.save(out_path)
    print(f"已保存: {out_path}")
    return score_sheet_name


# ------------------------- 内置语义校验 -------------------------

def eq(a, b):
    """容差比较:数值上相等即视为相同(57 与 57.0 显示无差别)"""
    if a == b:
        return True
    if isinstance(a, str) or isinstance(b, str):
        return False
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return False


def _load_matrix(path, sname=None, normal=False):
    """加载某个xlsx(或其指定sheet)为值矩阵。
    默认read_only流式（快）；normal=True时普通模式——read_only不解析合并单元格，
    含纵向合并的sheet(如学生明细)必须用普通模式，否则合并区读出假值导致比对失真。
    返回 {sheet名: rows}，rows为按行填充的二维列表（短行右侧补None）。"""
    wb = load_workbook(path, read_only=not normal, data_only=True)
    out = {}
    names = [sname] if sname else wb.sheetnames
    for nm in names:
        ws = wb[nm]
        if not normal and hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()   # 学校导出文件的dimension元数据常不可靠(只标A1)，强制全量流式扫描
        rows = []
        width = 0
        for tup in ws.iter_rows(values_only=True):
            row = list(tup)
            width = max(width, len(row))
            rows.append(row)
        for row in rows:                     # 补齐宽度，保证二维矩形
            if len(row) < width:
                row.extend([None] * (width - len(row)))
        out[nm] = rows
    wb.close()
    return out


def _cell(rows, r, c):
    """1-based取值，越界返回None"""
    if 1 <= r <= len(rows):
        row = rows[r - 1]
        if 1 <= c <= len(row):
            return row[c - 1]
    return None


def verify(raw_dir, out_path, score_sheet_name):
    classes = find_class_files(raw_dir)
    ok = True

    def check(cond, msg):
        nonlocal ok
        print(("  [PASS] " if cond else "  [FAIL] ") + msg)
        ok = ok and bool(cond)

    base_sheets = detail_sheets(classes[0][1])

    # 每个文件只加载一次——旧版同一明细文件被重复加载 N×sheet 次，是主要耗时；
    # 明细文件含纵向合并单元格，须普通模式(read_only不解析合并)
    out_m = _load_matrix(out_path)
    detail_m = {detail: _load_matrix(detail, normal=True) for _, detail, _ in classes}
    score_rows = {score: list(_load_matrix(score, wb_first_sheet(score)).values())[0]
                  for _, _, score in classes}
    tmap = extract_type_map(classes[0][1])

    check(list(out_m.keys()) == [score_sheet_name] + base_sheets,
          f"sheet清单顺序 {list(out_m.keys())}")

    # 得分汇总sheet: 总行数=各班之和;逐格比对(题号表头改写除外);表头改写正确性
    out_rows = out_m[score_sheet_name]
    total = 0
    for cls, _, score in classes:
        src = score_rows[score]
        nrow, ncol = len(src), len(src[0]) if src else 0
        s = total + 1
        mismatch = 0
        for r in range(1, nrow + 1):
            is_hdr = _cell(src, r, 1) == "序号"
            for c in range(1, ncol + 1):
                a = _cell(src, r, c)
                if is_hdr and as_int(a) is not None:  # 题号表头允许改写
                    continue
                if not eq(a, _cell(out_rows, s + r - 1, c)):
                    mismatch += 1
        check(mismatch == 0, f"{cls} 得分明细逐格比对({nrow}行×{ncol}列) 差异格数={mismatch}")
        # 表头改写检查:所有原为纯数字的表头都带上了题型前缀
        bad_hdr = []
        n_renamed = 0
        for c in range(2, ncol + 1):
            n = as_int(_cell(src, 2, c))
            if n is None:
                continue
            n_renamed += 1
            b = _cell(out_rows, s + 1, c)
            expect = f"{tmap[n]}{n}" if n in tmap else str(n)
            if b != expect:
                bad_hdr.append((c, n, b, expect))
        check(not bad_hdr, f"{cls} 表头题号改写正确({n_renamed}个题号列,异常:{bad_hdr[:3]})")
        total += nrow
    check(len(out_rows) == total, f"{score_sheet_name} 总行数 {len(out_rows)} = 各班之和 {total}")

    # 答题明细各sheet: 总行数=各班之和;逐格全量比对 + 首列标识集合
    for sname in base_sheets:
        out_rows = out_m[sname]
        total = 0
        for cls, detail, _ in classes:
            src = detail_m[detail][sname]
            nrow, ncol = len(src), len(src[0]) if src else 0
            ids_out, ids_src = set(), set()
            mismatch = 0
            for r in range(1, nrow + 1):
                for c in range(1, ncol + 1):
                    if not eq(_cell(src, r, c), _cell(out_rows, total + r, c)):
                        mismatch += 1
                v = _cell(src, r, 1)
                if r >= 3 and v is not None:
                    ids_src.add(str(as_int(v) if as_int(v) is not None else v))
                v2 = _cell(out_rows, total + r, 1)
                if r >= 3 and v2 is not None:
                    ids_out.add(str(as_int(v2) if as_int(v2) is not None else v2))
            check(mismatch == 0, f"{sname} {cls} 逐格全量比对 差异格数={mismatch}")
            check(ids_out == ids_src, f"{sname} {cls} 首列标识集合一致({len(ids_src)}个)")
            total += nrow
        check(len(out_rows) == total, f"{sname} 总行数 {len(out_rows)} = 各班之和 {total}")

    print("校验结果:", "全部通过" if ok else "存在失败项!!")
    return ok


def wb_first_sheet(path):
    """只取sheet名列表第一个（供得分明细单表加载）"""
    wb = load_workbook(path, read_only=True)
    nm = wb.sheetnames[0]
    wb.close()
    return nm


def main():
    ap = argparse.ArgumentParser(description="学测成绩文档四合一:答题明细+得分明细 → 总文档")
    ap.add_argument("raw_dir", help="原始文件夹(含各班答题明细与得分明细xlsx)")
    ap.add_argument("out_path", help="输出xlsx路径")
    ap.add_argument("--backup", help="输出文件已存在时,先备份到此路径", default=None)
    ap.add_argument("--score-sheet-name", help="得分汇总sheet名称(默认:源为班级小题明细→'班级小题得分',否则沿用源sheet名)",
                    default=None)
    args = ap.parse_args()

    if os.path.exists(args.out_path):
        if args.backup:
            shutil.move(args.out_path, args.backup)
            print(f"已将现有文件备份为: {args.backup}")
        else:
            raise SystemExit(f"错误: {args.out_path} 已存在(可用 --backup 指定备份路径)")

    score_sheet_name = build(args.raw_dir, args.out_path, args.score_sheet_name)
    if not verify(args.raw_dir, args.out_path, score_sheet_name):
        sys.exit(1)


if __name__ == "__main__":
    main()
